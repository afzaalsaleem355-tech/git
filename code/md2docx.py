# -*- coding: utf-8 -*-
"""把本项目 Markdown 初稿转换为 Word (.docx)。使用捆绑 python-docx。"""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

SRC = sys.argv[1] if len(sys.argv) > 1 else "docs/开题报告_阻燃分子生成式设计_初稿.md"
DST = sys.argv[2] if len(sys.argv) > 2 else SRC.rsplit(".", 1)[0] + ".docx"

def set_east_asia(run, font="宋体"):
    run.font.name = "Times New Roman"
    r = run._element.rPr
    if r is None:
        run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)

def add_inline(par, text, bold=False, code=False, size=None, color=None):
    # 解析 **粗体** 与 `行内代码`
    pattern = re.compile(r'(\*\*.+?\*\*|`[^`]+`)')
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            r = par.add_run(text[pos:m.start()]); _fmt(r, bold, code, size, color)
        tok = m.group(0)
        if tok.startswith("**"):
            r = par.add_run(tok[2:-2]); _fmt(r, True, code, size, color)
        else:
            r = par.add_run(tok[1:-1]); _fmt(r, bold, True, size, color)
        pos = m.end()
    if pos < len(text):
        r = par.add_run(text[pos:]); _fmt(r, bold, code, size, color)

def _fmt(run, bold, code, size, color):
    run.bold = bold
    if code:
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
    else:
        set_east_asia(run, "宋体")
    if size:
        run.font.size = Pt(size)
    if color and not code:
        run.font.color.rgb = color

def add_code_block(doc, lines):
    for ln in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(ln if ln else " ")
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x20, 0x20, 0x20)

doc = Document()
# 默认样式（中文）
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

lines = open(SRC, encoding="utf-8").read().splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    stripped = line.strip()

    if stripped.startswith("```"):
        # 代码块
        buf = []
        i += 1
        while i < len(lines) and not lines[i].strip().startswith("```"):
            buf.append(lines[i]); i += 1
        add_code_block(doc, buf)
        i += 1
        continue

    if stripped == "---":
        i += 1
        continue

    if stripped.startswith("> "):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.space_after = Pt(4)
        add_inline(p, stripped[2:])
        for r in p.runs:
            r.italic = True
            r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        i += 1
        continue

    if stripped.startswith("### "):
        p = doc.add_heading(level=3)
        add_inline(p, stripped[4:], size=13, color=RGBColor(0,0,0))
        for r in p.runs: set_east_asia(r, "黑体")
        i += 1
        continue

    if stripped.startswith("## "):
        p = doc.add_heading(level=2)
        add_inline(p, stripped[3:], size=15, color=RGBColor(0,0,0))
        for r in p.runs: set_east_asia(r, "黑体")
        i += 1
        continue

    if stripped.startswith("# "):
        p = doc.add_heading(level=0)
        add_inline(p, stripped[2:], size=20, color=RGBColor(0,0,0))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs: set_east_asia(r, "黑体")
        i += 1
        continue

    if stripped.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        add_inline(p, stripped[2:])
        i += 1
        continue

    # 表格：当前行含 | 且下一行为 |---|---| 形式
    if "|" in line and i + 1 < len(lines) and re.match(r'^\s*\|?[\s:\-|]+\|?\s*$', lines[i+1]) and "-" in lines[i+1]:
        table_lines = []
        while i < len(lines) and "|" in lines[i]:
            table_lines.append(lines[i]); i += 1
        header = [c.strip() for c in table_lines[0].strip().strip("|").split("|")]
        rows = []
        for ln in table_lines[2:]:
            rows.append([c.strip() for c in ln.strip().strip("|").split("|")])
        t = doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        hdr = t.rows[0].cells
        for j, htxt in enumerate(header):
            hdr[j].text = ""
            add_inline(hdr[j].paragraphs[0], htxt, bold=True)
        for row in rows:
            cells = t.add_row().cells
            for j in range(len(header)):
                cells[j].text = ""
                add_inline(cells[j].paragraphs[0], row[j] if j < len(row) else "")
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        continue

    # 普通段落（含数字列表，保留原文字序号）
    if stripped:
        p = doc.add_paragraph()
        if re.match(r'^\d+\.\s', stripped):
            p.paragraph_format.left_indent = Cm(0.5)
        add_inline(p, stripped)
    i += 1

doc.save(DST)
print("已生成:", DST)

# 回读校验
from docx import Document as D2
d = D2(DST)
print("段落数:", len(d.paragraphs), "| 表格数:", len(d.tables))
